import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Helper to apply RTL and Arabic formatting
    def set_rtl(p):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>'))

    def add_run(p, text, font_size=11, bold=False, color_rgb=(30, 41, 59), italic=False):
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(font_size)
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>')
        rPr.append(rFonts)
        return run

    # Header / Brand Title
    p_brand = doc.add_paragraph()
    set_rtl(p_brand)
    add_run(p_brand, "مصنع جولد كلين | Gold Clean", font_size=12, bold=True, color_rgb=(180, 83, 9))

    # Main Document Title
    p_title = doc.add_paragraph()
    set_rtl(p_title)
    add_run(p_title, "تقرير إداري: تعزيز الأمان وحماية البيانات وتسوية شحنات المتجر", font_size=20, bold=True, color_rgb=(15, 23, 42))

    # Subtitle
    p_sub = doc.add_paragraph()
    set_rtl(p_sub)
    add_run(p_sub, "تأمين العمليات التجارية، حماية الأرباح مع ميزة الشراء كضيف، وتحديث حالات الشحنات المعلقة", font_size=12, italic=True, color_rgb=(100, 116, 139))

    # Decorative Line
    p_line = doc.add_paragraph()
    set_rtl(p_line)
    p_line_run = p_line.add_run("―" * 55)
    p_line_run.font.color.rgb = RGBColor(217, 119, 6)
    p_line_run.bold = True

    # Metadata Box Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("تاريخ التقرير: ", "أكتوبر 2026"), ("إلى السادة: ", "مجلس الإدارة وإدارة العمليات")],
        [("إعداد: ", "الفريق التقني وإدارة المتجر"), ("حالة التنفيذ: ", "تم التأمين والتحديث بنجاح 100%")]
    ]

    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(3.3)
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            cell._tc.get_or_add_tcPr().append(shading)
            
            p = cell.paragraphs[0]
            set_rtl(p)
            label, val = meta_data[row_idx][col_idx]
            add_run(p, label, font_size=10, bold=True, color_rgb=(71, 85, 105))
            add_run(p, val, font_size=10, bold=(row_idx == 1 and col_idx == 1), 
                    color_rgb=(16, 149, 100) if (row_idx == 1 and col_idx == 1) else (15, 23, 42))

    doc.add_paragraph()

    # Section 1: Executive Summary & Context
    p_sec1 = doc.add_paragraph()
    set_rtl(p_sec1)
    add_run(p_sec1, "1. الملخص التنفيذي والسياق التجاري (لماذا قمنا بهذه التحديثات؟)", font_size=14, bold=True, color_rgb=(15, 23, 42))

    p_ctx1 = doc.add_paragraph()
    set_rtl(p_ctx1)
    add_run(p_ctx1, "انطلاقاً من الرؤية التسويقية لزيادة مبيعات المتجر ورفع معدل إتمام الطلبات (Conversion Rate)، تم اتخاذ قرار تجاري سليم للغاية وهو: ", font_size=11)
    add_run(p_ctx1, "إلغاء إلزامية تسجيل الدخول، وإتاحة الشراء السريع للعملاء كزائر أو ضيف (Guest Checkout) ", font_size=11, bold=True, color_rgb=(180, 83, 9))
    add_run(p_ctx1, "بدون اشتراط إنشاء حساب أو كتابة كلمات مرور معقدة.", font_size=11)

    # Highlight box for the cause
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = callout.rows[0].cells[0]
    c_cell.width = Inches(6.8)
    shading_c = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FEF3C7"/>')
    c_cell._tc.get_or_add_tcPr().append(shading_c)
    borders_c = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="single" w:sz="36" w:space="0" w:color="D97706"/>
        </w:tcBorders>
    ''')
    c_cell._tc.get_or_add_tcPr().append(borders_c)
    cp = c_cell.paragraphs[0]
    set_rtl(cp)
    add_run(cp, "💡 النقطة الجوهرية للإدارة:\n", font_size=11, bold=True, color_rgb=(146, 64, 14))
    add_run(cp, "في السابق، كان النظام القديم يعتمد على «حساب جوجل الموثق» لمعرفة هوية كل مشتري وتحديد ما يحق له فعله. "
                "وعندما قمنا بإلغاء اشتراط تسجيل الدخول لراحة العميل، انفتحت أبواب النظام لجمهور غير مسجل. "
                "وكان لزاماً علينا بناء «بوابات حماية ذكية» تفصل بين راحة العميل في الشراء السريع كزائر، "
                "وبين منع استغلال أي مجهول أو روبوت آلي لهذا الباب المفتوح. "
                "كافة النقاط الأمنية التي تم علاجها هي الاستجابة التقنية المباشرة لضمان نجاح قرار الشراء بدون تسجيل بأمان تام.", 
                font_size=10.5, color_rgb=(120, 53, 15))

    doc.add_paragraph()

    # Section 2: Key Risks & Solutions
    p_sec2 = doc.add_paragraph()
    set_rtl(p_sec2)
    add_run(p_sec2, "2. أبرز المخاطر المحتملة وكيف تمت معالجتها لحماية المتجر", font_size=14, bold=True, color_rgb=(15, 23, 42))

    risks = [
        {
            "num": "أ",
            "title": "حماية أسعار المنتجات من التلاعب (تثبيت الأسعار للمدير فقط)",
            "why": "للسماح للزوار بتقييم المنتجات ووضع النجوم بحرية، كانت قاعدة البيانات تسمح بتحديث بيانات المنتج.",
            "risk": "احتمالية قيام مستخدم متقدم بالتلاعب بسعر المنظفات وجعلها بقيمة صفر أو جنيه واحد.",
            "solution": "تم إغلاق التعديل بالكامل وحصره حصرياً في إدارة المتجر المعتمدة. كما تمت برمجة استثناء ذكي يسمح للعميل فقط بإرسال تقييم النجوم (من 1 إلى 5) دون أي قدرة على لمس السعر أو الوصف."
        },
        {
            "num": "ب",
            "title": "حماية سرية بيانات العملاء وعناوينهم (منع التسريب)",
            "why": "لتسهيل تتبع طلب العميل كضيف بعد الشراء، كان النظام يتعامل مع طلبات الضيوف بصيغة مفتوحة القراءة.",
            "risk": "احتمالية قيام أي منافس أو جهة خارجية بسحب قائمة بأسماء كل عملاء المتجر وأرقام هواتفهم وعناوين منازلهم دفعة واحدة.",
            "solution": "تم حظر استخراج أو تحميل طلبات العملاء بشكل جماعي نهائياً. وفي نفس الوقت، يحتفظ العميل وحده برقم طلبه السري على جهازه ليتمكن من متابعة شحنته دون أن يرى أي شخص طلبات الآخرين."
        },
        {
            "num": "ج",
            "title": "حماية ميزانية الشحن ومنع الطلبات الوهمية مع شركة J&T Express",
            "why": "المتجر مربوط إلكترونياً وبشكل فوري مع شركة الشحن لإصدار البوليصة لحظة تأكيد الطلب.",
            "risk": "قيام روبوتات إنترنت مجهولة بإرسال مئات الطلبات الوهمية بأرقام خاطئة، مما ينشئ بوالص حقيقية تكبد الشركة تكاليف شحن أو غرامات إلغاء مع شركة J&T.",
            "solution": "تم تركيب نظامين للحماية: (1) فحص آلي صارم لأرقام الهواتف والتأكد من أنها أرقام مصرية صحيحة (11 رقماً تبدأ بـ 010 أو 011 أو 012 أو 015). (2) وضع حد أقصى للطلبات من نفس الجهاز لمنع الإغراق الآلي (Rate Limiting)."
        },
        {
            "num": "د",
            "title": "حماية أدوات الذكاء الاصطناعي من الاستنزاف (AI SEO Protection)",
            "why": "يحتوي المتجر على ميزة متطورة تستخدم الذكاء الاصطناعي لكتابة أوصاف المنتجات لمحركات البحث.",
            "risk": "استهلاك رصيد الشركة في خدمات الذكاء الاصطناعي من قِبل أطراف خارجية بدون وجه حق.",
            "solution": "تم غلق وتشفير هذه الميزة لتكون حصرية فقط لحسابات مدراء المتجر الموثقين عبر Google."
        }
    ]

    for item in risks:
        p_item = doc.add_paragraph()
        set_rtl(p_item)
        add_run(p_item, f"({item['num']}) {item['title']}\n", font_size=12, bold=True, color_rgb=(30, 58, 138))
        
        p_desc = doc.add_paragraph()
        set_rtl(p_desc)
        add_run(p_desc, "• سبب الظهور: ", font_size=10.5, bold=True, color_rgb=(100, 116, 139))
        add_run(p_desc, item['why'] + "\n", font_size=10.5)
        
        add_run(p_desc, "• الخطر السابق: ", font_size=10.5, bold=True, color_rgb=(225, 29, 72))
        add_run(p_desc, item['risk'] + "\n", font_size=10.5)
        
        add_run(p_desc, "• الإجراء المنفذ: ", font_size=10.5, bold=True, color_rgb=(16, 149, 100))
        add_run(p_desc, item['solution'], font_size=10.5, bold=True, color_rgb=(15, 23, 42))

    doc.add_paragraph()

    # Section 3: Summary Comparison Table
    p_sec3 = doc.add_paragraph()
    set_rtl(p_sec3)
    add_run(p_sec3, "3. جدول المقارنة الإداري (قبل وبعد تطبيق الحماية)", font_size=14, bold=True, color_rgb=(15, 23, 42))

    tbl = doc.add_table(rows=5, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    headers = ["المجال التجاري والأمني", "الوضع السابق", "الوضع الحالي بعد الحماية المعتمدة"]
    table_widths = [Inches(1.8), Inches(2.5), Inches(2.5)]

    for c_idx, h_text in enumerate(headers):
        cell = tbl.rows[0].cells[c_idx]
        cell.width = table_widths[c_idx]
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        add_run(p, h_text, font_size=10.5, bold=True, color_rgb=(255, 255, 255))

    rows_data = [
        ("أسعار المنتجات", "عرضة للتعديل من أي متصفح خارجي", "محمية ومقفلة بنسبة 100% للمدير فقط مع بقاء ميزة تقييم الزوار"),
        ("خصوصية العملاء والطلبات", "إمكانية تحميل كافة طلبات الضيوف وهواتفهم", "عزل تام، لا أحد يرى سوى طلبه الشخصي برقم الطلب السري"),
        ("حساب شركة الشحن (J&T)", "عرضة للإغراق بالطلبات الوهمية وتوليد بوالص عشوائية", "فلترة صارمة لأرقام الهواتف المصرية وتحديد سقف المحاولات"),
        ("رصيد الذكاء الاصطناعي", "متاح للاستهلاك المفتوح", "محمي ومخصص حصراً لمدراء ومسؤولي المتجر")
    ]

    for r_idx, r_data in enumerate(rows_data):
        row = tbl.rows[r_idx + 1]
        for c_idx, text in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = table_widths[c_idx]
            bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            set_rtl(p)
            is_good = (c_idx == 2)
            is_bad = (c_idx == 1)
            add_run(p, text, font_size=10, bold=(c_idx == 0 or is_good),
                    color_rgb=(16, 149, 100) if is_good else ((185, 28, 28) if is_bad else (15, 23, 42)))

    doc.add_paragraph()

    # Section 4: Operational Update (Manual status updates for un-synced waybills)
    p_sec4 = doc.add_paragraph()
    set_rtl(p_sec4)
    add_run(p_sec4, "4. تقرير تشغيلي: تسوية وتحديث حالات الشحنات المعلقة يدوياً في المتجر", font_size=14, bold=True, color_rgb=(15, 23, 42))

    p_op_desc = doc.add_paragraph()
    set_rtl(p_op_desc)
    add_run(p_op_desc, "نظراً لأن بعض الشحنات قد تواجه تأخيراً في وصول التحديث الآلي لحالاتها من نظام شركة الشحن (J&T Express) إلى الموقع الإلكتروني (نتيجة تأخر مزامنة أجهزة المناديب أو أنظمة التتبع)، "
                       "وحرصاً منا على دقة الإحصائيات والتقارير المالية أمام الإدارة، ", font_size=10.5)
    add_run(p_op_desc, "قام مسؤول المتجر بعمل مطابقة ميدانية وتحديث يدوي لحالات 17 شحنة معلقة ", font_size=10.5, bold=True, color_rgb=(180, 83, 9))
    add_run(p_op_desc, "على قاعدة بيانات المتجر لتتطابق تماماً مع الواقع الفعلي للتحصيل والتسليم.", font_size=10.5)

    # Operational metrics boxes table
    op_metrics_table = doc.add_table(rows=1, cols=3)
    op_metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    op_metrics_table.autofit = False

    metrics_info = [
        ("إجمالي الشحنات المسواة", "17 شحنة", "F1F5F9", (30, 41, 59)),
        ("شحنات تم تسليمها بنجاح (تم الاستلام)", "16 شحنة", "ECFDF5", (16, 149, 100)),
        ("شحنات مرتجعة للمخزن (مرتجع)", "شحنة واحدة", "FFF1F2", (225, 29, 72))
    ]

    for m_idx, (m_label, m_val, m_bg, m_color) in enumerate(metrics_info):
        m_cell = op_metrics_table.rows[0].cells[m_idx]
        m_cell.width = Inches(2.26)
        shd_m = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{m_bg}"/>')
        m_cell._tc.get_or_add_tcPr().append(shd_m)
        mp = m_cell.paragraphs[0]
        set_rtl(mp)
        add_run(mp, m_label + "\n", font_size=9.5, bold=True, color_rgb=(71, 85, 105))
        add_run(mp, m_val, font_size=13, bold=True, color_rgb=m_color)

    doc.add_paragraph()

    # Table of Updated Waybills
    p_tbl_lbl = doc.add_paragraph()
    set_rtl(p_tbl_lbl)
    add_run(p_tbl_lbl, "📋 تفاصيل أرقام البوالص والحالات المحدثة:", font_size=11, bold=True, color_rgb=(30, 58, 138))

    delivered_waybills = [
        "JEG000546380797", "JEG000550087130", "JEG000549692956", "JEG000549881135",
        "JEG000546137639", "JEG000538006259", "JEG000549213258", "JEG000545482954",
        "JEG000549644995", "JEG000544848213", "JEG000548922585", "JEG000541587061",
        "JEG000545055682", "JEG000544332175", "JEG000545355717", "JEG000546713070"
    ]
    returned_waybills = ["JEG000549071358"]

    all_wb_data = []
    for idx, wb in enumerate(delivered_waybills, start=1):
        all_wb_data.append((str(idx), wb, "تم الاستلام ✅", "تم تسليم الشحنة للعميل بنجاح وتأكيد التحصيل"))
    for idx, wb in enumerate(returned_waybills, start=len(delivered_waybills)+1):
        all_wb_data.append((str(idx), wb, "مرتجع 🔄", "تعذر تسليم الشحنة وارتجاع البضاعة للمخزن"))

    wb_table = doc.add_table(rows=len(all_wb_data) + 1, cols=4)
    wb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    wb_table.autofit = False

    wb_widths = [Inches(0.6), Inches(2.2), Inches(1.5), Inches(2.5)]
    wb_headers = ["م", "رقم بوليصة الشحن (Waybill)", "الحالة المحدثة", "الملاحظة التشغيلية"]

    for c_idx, h_text in enumerate(wb_headers):
        cell = wb_table.rows[0].cells[c_idx]
        cell.width = wb_widths[c_idx]
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        add_run(p, h_text, font_size=9.5, bold=True, color_rgb=(255, 255, 255))

    for r_idx, row_item in enumerate(all_wb_data):
        row = wb_table.rows[r_idx + 1]
        is_returned = ("مرتجع" in row_item[2])
        row_bg = "FFF1F2" if is_returned else ("FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
        for c_idx, val in enumerate(row_item):
            cell = row.cells[c_idx]
            cell.width = wb_widths[c_idx]
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{row_bg}"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            set_rtl(p)
            color = (225, 29, 72) if is_returned else ((16, 149, 100) if c_idx == 2 else (30, 41, 59))
            add_run(p, val, font_size=9, bold=(c_idx in [1, 2]), color_rgb=color)

    doc.add_paragraph()

    # Section 5: Business Value & Recommendations
    p_sec5 = doc.add_paragraph()
    set_rtl(p_sec5)
    add_run(p_sec5, "5. العائد التجاري المباشر وتوصيات الخطوة القادمة", font_size=14, bold=True, color_rgb=(15, 23, 42))

    points = [
        ("الحفاظ على تجربة شراء سريعة وبدون تعقيد:", " لم يتأثر العميل إطلاقاً بهذه التحديثات الأمنية؛ فهو لا يزال يطلب منتجاته بنقرة واحدة كزائر وبدون إنشاء حساب."),
        ("حماية سمعة مصنع Gold Clean في السوق:", " حماية أرقام هواتف وعناوين العملاء التزام أخلاقي وتجاري يمنع أي تسريب للمنافسين."),
        ("حماية الميزانية والأرباح:", " إغلاق الباب أمام التلاعب بأسعار المنظفات أو تحميل المتجر تكاليف بوالص شحن وهمية مع J&T Express."),
        ("دقة البيانات المالية وإحصائيات التسليم:", " مطابقة وتحديث حالات الشحنات الـ 17 يضمن مطابقة مبيعات الموقع مع كشوفات حساب شركة الشحن وأرقام التحصيل الفعلية."),
        ("التوصية القادمة لزيادة التحصيل:", " نقترح تفعيل خاصية «تأكيد الطلب الفوري عبر واتساب»، بحيث يرسل النظام تفاصيل الشحنة للعميل بنقرة واحدة لتأكيد جديته قبل خروج المندوب، مما يخفض نسبة المرتجعات (RTO) بشكل قياسي.")
    ]

    for title, desc in points:
        p_pt = doc.add_paragraph()
        set_rtl(p_pt)
        add_run(p_pt, "✔ " + title, font_size=10.5, bold=True, color_rgb=(180, 83, 9))
        add_run(p_pt, desc, font_size=10.5, color_rgb=(30, 41, 59))

    doc.add_paragraph()

    # Sign-off Box
    p_sign = doc.add_paragraph()
    set_rtl(p_sign)
    add_run(p_sign, "خاتمة وتأكيد جاهزية النظام:\n", font_size=11, bold=True, color_rgb=(15, 23, 42))
    add_run(p_sign, "تم اختبار جميع الإجراءات المذكورة في بيئة العمل الحية والتأكد من أنها تعمل بكفاءة تامة وتوفر الحماية القصوى للمتجر مع استمرار استقبال الطلبات بكل سلاسة ودقة البيانات التشغيلية.\n\n", font_size=10.5)
    add_run(p_sign, "فريق التطوير والتشغيل التقني | متجر Gold Clean الإلكتروني", font_size=10.5, bold=True, color_rgb=(100, 116, 139))

    filename = "Executive_Security_Report_GoldClean.docx"
    doc.save(filename)
    print(f"Report generated successfully: {filename}")

if __name__ == '__main__':
    create_report()
